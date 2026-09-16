# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

DST = r'C:\Users\30216\Desktop\个人财务管理与可视化看板\deliverables\个人财务管理与可视化看板-综合实训设计报告.docx'
doc = Document(DST)
print('=== 含“年”的段落 ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if '年' in t and ('二' in t or '2026' in t):
        print(i, repr(t), '| runs:', [r.text for r in p.runs])
print()
print('=== 封面姓名单元格 ===')
t0 = doc.tables[0]
print('学生姓名:', repr(t0.cell(3,1).text))
print('学号:', repr(t0.cell(4,1).text))

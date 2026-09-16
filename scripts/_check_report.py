# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

DST = r'C:\Users\30216\Desktop\个人财务管理与可视化看板\deliverables\个人财务管理与可视化看板-综合实训设计报告.docx'
doc = Document(DST)

print('=== 段落总数:', len(doc.paragraphs), ' 表格数:', len(doc.tables))
print()
print('=== 表格检查 ===')
for ti, tb in enumerate(doc.tables):
    print(f'-- 表格{ti}: {len(tb.rows)}x{len(tb.columns)} --')
    for row in tb.rows[:8]:
        cells = []
        seen = set()
        for c in row.cells:
            if id(c._tc) not in seen:
                seen.add(id(c._tc))
                cells.append(c.text.strip().replace('\n', '/')[:24])
        print('   ', ' | '.join(cells))
print()
print('=== 正文标题/图片分布 ===')
n_img = 0
for p in doc.paragraphs:
    has_img = p._p.findall('.//' + qn('a:blip'))
    if has_img:
        n_img += 1
        print(f'  [图片] 段落: {p.text[:40]}')
for p in doc.paragraphs:
    t = p.text.strip()
    if t and (t.startswith(('一 ', '二 ', '三 ', '四 ', '五 ', '附 ')) or t.startswith('（') or t.startswith(('1.', '2.', '3.', '4.'))):
        print('  [标题]', t)
print('图片总数:', n_img)
